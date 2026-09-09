from pathlib import Path

import pymupdf  # PyMuPDF
from docx import Document


def extract_from_pdf(file_path: str) -> str:
    """Extract text from a PDF file."""

    text = []

    try:
        pdf = pymupdf.open(file_path)

        for page in pdf:
            page_text = page.get_text("text")

            if page_text.strip():
                text.append(page_text)

        pdf.close()

    except Exception as e:
        raise RuntimeError(
            f"Failed to extract text from PDF '{file_path}': {e}"
        )

    return "\n".join(text).strip()


def extract_from_docx(file_path: str) -> str:
    """Extract text from a DOCX file."""

    try:
        document = Document(file_path)

        paragraphs = [
            paragraph.text.strip()
            for paragraph in document.paragraphs
            if paragraph.text.strip()
        ]

        # Extract text from tables as well
        for table in document.tables:
            for row in table.rows:
                row_text = " | ".join(
                    cell.text.strip()
                    for cell in row.cells
                    if cell.text.strip()
                )

                if row_text:
                    paragraphs.append(row_text)

    except Exception as e:
        raise RuntimeError(
            f"Failed to extract text from DOCX '{file_path}': {e}"
        )

    return "\n".join(paragraphs).strip()


def extract_from_txt(file_path: str) -> str:
    """Extract text from a TXT file."""

    try:
        return Path(file_path).read_text(
            encoding="utf-8",
            errors="ignore"
        ).strip()

    except Exception as e:
        raise RuntimeError(
            f"Failed to extract text from TXT '{file_path}': {e}"
        )


def extract_text(file_path: str) -> str:
    """
    Automatically detect the document type
    and extract its text.
    """

    path = Path(file_path)

    if not path.exists():
        raise FileNotFoundError(
            f"File not found: {file_path}"
        )

    extension = path.suffix.lower()

    if extension == ".pdf":
        return extract_from_pdf(file_path)

    if extension == ".docx":
        return extract_from_docx(file_path)

    if extension == ".txt":
        return extract_from_txt(file_path)

    raise ValueError(
        f"Unsupported file format: {extension}"
    )